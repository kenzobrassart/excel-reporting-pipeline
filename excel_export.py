import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage


# couleurs du rapport - palette orange (codes hexa sans le #, openpyxl veut ce format la)
ORANGE_FONCE = "8B3A1A"   # marron-orange pour les en-tetes principales
ORANGE_CLAIR = "FDE8D8"   # peche tres clair pour les zebrures
ORANGE = "E67E22"         # orange vif pour les valeurs positives
ROUGE = "C0392B"          # rouge profond pour les valeurs negatives
BLANC = "FFFFFF"

# nuances intermediaires utilisees pour les sous-titres et entetes secondaires
ORANGE_MOYEN = "B8460E"   # orange-brun moyen
ORANGE_LIGHT = "D67D2E"   # orange un peu plus clair


# fonctions utilitaires de mise en forme, pour pas avoir a tout retaper a chaque fois

def style_entete(cell, bg=ORANGE_FONCE, taille=11):
    # style des cellules d'en-tete : gras, blanc, fond colore, centre
    cell.font = Font(bold=True, color=BLANC, size=taille, name="Arial")
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_cellule(cell, bold=False, taille=10, bg=None, couleur="000000", centre=False):
    # style des cellules de donnees, avec plein d'options optionnelles
    cell.font = Font(bold=bold, size=taille, name="Arial", color=couleur)
    cell.alignment = Alignment(horizontal="center" if centre else "left", vertical="center")
    if bg:
        cell.fill = PatternFill("solid", start_color=bg)


def bordure_fine():
    # petite bordure grise sur les 4 cotes pour delimiter les cellules
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def ajuster_colonnes(ws, mini=12, maxi=40):
    # ajuste la largeur de chaque colonne en fonction du contenu le plus long
    # (avec un mini et un maxi pour pas avoir des colonnes ridicules ou enormes)
    for col in ws.columns:
        longueur = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min(longueur + 3, maxi), mini)


def ecrire_tcd(ws, tcd, ligne_titre, titre, n_colonnes_pivot, note=None):
    # fonction generique pour ecrire un TCD dans la feuille
    # comme ca on evite de copier coller le code 3 fois
    # n_colonnes_pivot = nombre de modalites de la variable en colonne (ex: 3 pour les classes)

    # titre du TCD sur toute la largeur
    nb_total = 1 + 3 * (n_colonnes_pivot + 1)  # 1 pour la 1ere colonne + 3 blocs (mean/sum/count) avec totaux
    derniere_lettre = get_column_letter(nb_total)
    ws.merge_cells(f"A{ligne_titre}:{derniere_lettre}{ligne_titre}")
    ws[f"A{ligne_titre}"] = titre
    style_entete(ws[f"A{ligne_titre}"], taille=12)
    ws.row_dimensions[ligne_titre].height = 26

    # entetes des colonnes du TCD
    cols_pivot = list(tcd.columns)
    entetes = ["Categorie"] + [f"Survie moy {c} (%)" if "mean" in c else
                                f"Nb survivants {c.split('_')[1]}" if "sum" in c else
                                f"Nb passagers {c.split('_')[1]}"
                                for c in cols_pivot]
    ligne_entete = ligne_titre + 2
    for j, h in enumerate(entetes, start=1):
        cell = ws.cell(row=ligne_entete, column=j, value=h)
        style_entete(cell, bg=ORANGE_MOYEN, taille=9)
        cell.border = bordure_fine()

    # remplissage des lignes du TCD
    for i, (idx, row) in enumerate(tcd.iterrows(), start=ligne_entete + 1):
        bg = ORANGE_CLAIR if i % 2 == 0 else BLANC
        ws.cell(row=i, column=1, value=str(idx).capitalize())
        style_cellule(ws.cell(row=i, column=1), bold=True, bg=bg)
        ws.cell(row=i, column=1).border = bordure_fine()

        for j, (col_nom, val) in enumerate(zip(cols_pivot, row.values), start=2):
            if "mean" in col_nom:
                texte = f"{val*100:.1f}%" if pd.notna(val) else "-"
            else:
                texte = str(int(val)) if pd.notna(val) else "-"
            cell = ws.cell(row=i, column=j, value=texte)
            style_cellule(cell, bg=bg, centre=True)
            cell.border = bordure_fine()

    # note optionnelle en dessous du TCD (pour signaler les pieges statistiques)
    derniere_ligne = ligne_entete + 1 + len(tcd)
    if note:
        ws.cell(row=derniere_ligne + 1, column=1, value=note)
        style_cellule(ws.cell(row=derniere_ligne + 1, column=1), taille=9, couleur="666666")
        derniere_ligne += 1

    return derniere_ligne


def faire_tcd(df, index_col):
    # cree un TCD avec mean/sum/count, totaux marginaux, colonnes aplaties
    tcd = df.pivot_table(
        values="survived",
        index=index_col,
        columns="class",
        aggfunc=["mean", "sum", "count"],
        margins=True,
        margins_name="Total",
        observed=True
    )
    tcd.columns = [f"{agg}_{col}" for agg, col in tcd.columns]
    return tcd


def faire_tcd_isalone(df):
    # TCD specifique : IsAlone x sexe (pas par classe ici)
    tcd = df.pivot_table(
        values="survived",
        index="IsAlone",
        columns="sex",
        aggfunc=["mean", "sum", "count"],
        margins=True,
        margins_name="Total"
    )
    tcd.columns = [f"{agg}_{col}" for agg, col in tcd.columns]
    # on remplace 0 et 1 par des libelles plus clairs
    tcd.index = tcd.index.map({0: "En famille", 1: "Seul", "Total": "Total"})
    return tcd


# Feuille 1 : donnees brutes nettoyees

def feuille_donnees(wb, df):
    ws = wb.active
    ws.title = "Donnees"

    # gros titre fusionne en haut sur toute la largeur du tableau
    ws.merge_cells("A1:O1")
    ws["A1"] = "Dataset Titanic - donnees nettoyees"
    style_entete(ws["A1"], taille=13)
    ws.row_dimensions[1].height = 28

    # ligne d'entetes : on reprend les noms des colonnes du dataframe
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=2, column=j, value=col)
        style_entete(cell, taille=10)
        cell.border = bordure_fine()

    # remplissage des lignes du dataframe, avec une couleur sur 2 (zebrures)
    # ca rend le tableau beaucoup plus lisible quand y a 891 lignes
    for i, (_, row) in enumerate(df.iterrows(), start=3):
        bg = ORANGE_CLAIR if i % 2 == 0 else BLANC
        for j, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=val)
            style_cellule(cell, bg=bg)
            cell.border = bordure_fine()

    ajuster_colonnes(ws)
    # fige les 2 premieres lignes (titre + entetes) pour qu'elles restent visibles
    # quand on scrolle vers le bas
    ws.freeze_panes = "A3"


# Feuille 2 : indicateurs cles

def feuille_indicateurs(wb, resultats):
    ws = wb.create_sheet("Indicateurs")

    ws.merge_cells("A1:D1")
    ws["A1"] = "Indicateurs cles"
    style_entete(ws["A1"], taille=13)
    ws.row_dimensions[1].height = 28

    # on construit la feuille bloc par bloc, en gardant un compteur de ligne
    # qui s'incremente au fur et a mesure
    ligne = 3

    # bloc 1 : taux global
    ws.merge_cells(f"A{ligne}:D{ligne}")
    ws[f"A{ligne}"] = "Taux de survie global"
    style_entete(ws[f"A{ligne}"], bg=ORANGE_MOYEN)
    ligne += 1
    ws[f"A{ligne}"] = "Indicateur"
    ws[f"B{ligne}"] = "Valeur"
    for col in ["A", "B"]:
        style_entete(ws[col + str(ligne)], bg=ORANGE_LIGHT, taille=10)
    ligne += 1
    ws[f"A{ligne}"] = "Taux de survie global"
    ws[f"B{ligne}"] = f"{resultats['taux_global']:.1f}%"
    style_cellule(ws[f"A{ligne}"])
    style_cellule(ws[f"B{ligne}"], bold=True, couleur=ORANGE, centre=True)
    for col in ["A", "B"]:
        ws[f"{col}{ligne}"].border = bordure_fine()
    ligne += 2

    # bloc 2 : par sexe
    ws.merge_cells(f"A{ligne}:D{ligne}")
    ws[f"A{ligne}"] = "Taux de survie par sexe"
    style_entete(ws[f"A{ligne}"], bg=ORANGE_MOYEN)
    ligne += 1
    ws[f"A{ligne}"] = "Sexe"
    ws[f"B{ligne}"] = "Taux (%)"
    for col in ["A", "B"]:
        style_entete(ws[col + str(ligne)], bg=ORANGE_LIGHT, taille=10)
    ligne += 1
    for sexe, taux in resultats["taux_par_sexe"].items():
        ws[f"A{ligne}"] = sexe.capitalize()
        ws[f"B{ligne}"] = f"{taux:.1f}%"
        # orange si > 50%, rouge sinon : repere visuel rapide
        couleur = ORANGE if taux > 50 else ROUGE
        style_cellule(ws[f"A{ligne}"])
        style_cellule(ws[f"B{ligne}"], bold=True, couleur=couleur, centre=True)
        for col in ["A", "B"]:
            ws[f"{col}{ligne}"].border = bordure_fine()
        ligne += 1
    ligne += 1

    # bloc 3 : par classe
    ws.merge_cells(f"A{ligne}:D{ligne}")
    ws[f"A{ligne}"] = "Taux de survie par classe"
    style_entete(ws[f"A{ligne}"], bg=ORANGE_MOYEN)
    ligne += 1
    ws[f"A{ligne}"] = "Classe"
    ws[f"B{ligne}"] = "Taux (%)"
    for col in ["A", "B"]:
        style_entete(ws[col + str(ligne)], bg=ORANGE_LIGHT, taille=10)
    ligne += 1
    for classe, taux in resultats["taux_par_classe"].items():
        ws[f"A{ligne}"] = str(classe)
        ws[f"B{ligne}"] = f"{taux:.1f}%"
        # seuil un peu plus bas ici (40% au lieu de 50%) car la 3eme classe est defavorisee
        couleur = ORANGE if taux > 40 else ROUGE
        style_cellule(ws[f"A{ligne}"])
        style_cellule(ws[f"B{ligne}"], bold=True, couleur=couleur, centre=True)
        for col in ["A", "B"]:
            ws[f"{col}{ligne}"].border = bordure_fine()
        ligne += 1
    ligne += 1

    # bloc 4 : age moyen et median
    ws.merge_cells(f"A{ligne}:D{ligne}")
    ws[f"A{ligne}"] = "Age moyen et median (survivants vs non-survivants)"
    style_entete(ws[f"A{ligne}"], bg=ORANGE_MOYEN)
    ligne += 1
    ws[f"A{ligne}"] = "Groupe"
    ws[f"B{ligne}"] = "Age moyen"
    ws[f"C{ligne}"] = "Age median"
    for col in ["A", "B", "C"]:
        style_entete(ws[col + str(ligne)], bg=ORANGE_LIGHT, taille=10)
    ligne += 1
    for label, moy, med in [
        ("Non-survivants",
         resultats["age_median"]["moyenne_non_survivants"],
         resultats["age_median"]["mediane_non_survivants"]),
        ("Survivants",
         resultats["age_median"]["moyenne_survivants"],
         resultats["age_median"]["mediane_survivants"])
    ]:
        ws[f"A{ligne}"] = label
        ws[f"B{ligne}"] = f"{moy} ans"
        ws[f"C{ligne}"] = f"{med} ans"
        style_cellule(ws[f"A{ligne}"])
        style_cellule(ws[f"B{ligne}"], centre=True)
        style_cellule(ws[f"C{ligne}"], centre=True)
        for col in ["A", "B", "C"]:
            ws[f"{col}{ligne}"].border = bordure_fine()
        ligne += 1

    ajuster_colonnes(ws, mini=20)


# Feuille 3 : tableau croise + 3 TCD

def feuille_tableau_croise(wb, pivot, df):
    ws = wb.create_sheet("Tableau croise")

    ws.merge_cells("A1:E1")
    ws["A1"] = "Taux de survie (%) par sexe et classe"
    style_entete(ws["A1"], taille=13)
    ws.row_dimensions[1].height = 28

    # premier tableau : juste les pourcentages de survie pour chaque combinaison
    # sexe x classe (ex : femme en 1ere classe = 96.8% de survie)
    ws.cell(row=3, column=1, value="Sexe / Classe")
    style_entete(ws.cell(row=3, column=1), bg=ORANGE_MOYEN)

    for j, col in enumerate(pivot.columns, start=2):
        cell = ws.cell(row=3, column=j, value=str(col))
        style_entete(cell, bg=ORANGE_LIGHT, taille=10)

    for i, (idx, row) in enumerate(pivot.iterrows(), start=4):
        bg = ORANGE_CLAIR if i % 2 == 0 else BLANC
        ws.cell(row=i, column=1, value=str(idx).capitalize())
        style_cellule(ws.cell(row=i, column=1), bold=True, bg=bg)
        ws.cell(row=i, column=1).border = bordure_fine()

        for j, val in enumerate(row, start=2):
            cell = ws.cell(row=i, column=j, value=f"{val:.1f}%")
            # code couleur en 3 niveaux : orange > 50%, rouge < 25%, noir entre les deux
            couleur = ORANGE if val > 50 else (ROUGE if val < 25 else "000000")
            style_cellule(cell, bold=(val > 50), couleur=couleur, bg=bg, centre=True)
            cell.border = bordure_fine()

    # ============================================================
    # TCD n°1 : sexe x classe (effectifs detailles)
    # ============================================================
    # un TCD = tableau croise dynamique = on croise 2 variables et on calcule
    # plusieurs stats : la moyenne du taux de survie, le nb de survivants,
    # le nb total de passagers, avec des totaux generaux dans la derniere ligne/colonne
    ligne_courante = 4 + len(pivot) + 2

    tcd1 = faire_tcd(df, "sex")
    ligne_courante = ecrire_tcd(
        ws, tcd1, ligne_courante,
        "TCD n°1 : Survie par sexe et classe (effectifs detailles)",
        n_colonnes_pivot=3
    )

    # ============================================================
    # TCD n°2 : tranche d'age x classe
    # ============================================================
    # ATTENTION : la tranche "senior" ne contient que 22 personnes au total
    # (14 en 1ere classe, 3 en 2eme, 5 en 3eme), donc les pourcentages
    # affiches sont a interpreter avec prudence pour cette ligne
    # (un seul senior qui change de statut fait varier le taux de plusieurs %)
    ligne_courante += 3
    tcd2 = faire_tcd(df, "age_cat")
    note_age = ("Note : la categorie 'senior' ne contient que 22 personnes au total. "
                "Les taux affiches pour cette ligne sont peu fiables statistiquement.")
    ligne_courante = ecrire_tcd(
        ws, tcd2, ligne_courante,
        "TCD n°2 : Survie par tranche d'age et classe",
        n_colonnes_pivot=3,
        note=note_age
    )

    # ============================================================
    # TCD n°3 : IsAlone x sexe
    # ============================================================
    # voyager seul ou en famille change-t-il les chances de survie ?
    # ici on croise IsAlone (notre variable indicatrice) avec le sexe
    # pas de probleme d'effectifs : la plus petite case contient 126 personnes
    ligne_courante += 3
    tcd3 = faire_tcd_isalone(df)
    ligne_courante = ecrire_tcd(
        ws, tcd3, ligne_courante,
        "TCD n°3 : Survie selon le fait de voyager seul, par sexe",
        n_colonnes_pivot=2  # 2 modalites de sexe (female, male)
    )

    ajuster_colonnes(ws, mini=15)


# Feuille 4 : graphiques

def feuille_graphiques(wb, chemins):
    ws = wb.create_sheet("Graphiques")

    ws.merge_cells("A1:X1")
    ws["A1"] = "Visualisations"
    style_entete(ws["A1"], taille=13)
    ws.row_dimensions[1].height = 28

    # on place les 3 images png cote a cote sur la meme ligne
    # chaque image fait 420px de large, donc on espace de ~8 colonnes a chaque fois
    positions = [
        ("survie_sexe", "A3"),
        ("survie_classe", "I3"),
        ("ages", "Q3"),
    ]

    for cle, ancre in positions:
        chemin = chemins.get(cle)
        if chemin and os.path.exists(chemin):
            img = XLImage(chemin)
            img.width = 420
            img.height = 300
            ws.add_image(img, ancre)


# fonction principale qui appelle tout dans l'ordre

def generer_rapport_excel(df, resultats, chemins, sortie="output/rapport_titanic.xlsx"):
    # cree le dossier de sortie si il existe pas deja
    os.makedirs(os.path.dirname(sortie), exist_ok=True)
    wb = Workbook()

    print("Construction du rapport excel")
    feuille_donnees(wb, df)
    feuille_indicateurs(wb, resultats)
    feuille_tableau_croise(wb, resultats["tableau_croise"], df)
    feuille_graphiques(wb, chemins)

    wb.save(sortie)
    print("Rapport sauvegarde :", sortie)
    return sortie
